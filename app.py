import streamlit as st
import json
from google import genai
from docx import Document
from pptx import Presentation
import io

# Setup trang Streamlit
st.set_page_config(page_title="AI Hỗ Trợ Giảng Dạy GDPT 2018", layout="wide")
st.title("Ứng Dụng Xuất Bản Giáo Án, Slide & Đề Thi (theo Công văn 7991)")

# Sidebar nhập API Key và Thông tin bài học
st.sidebar.header("⚙️ Cấu hình")
api_key = st.sidebar.text_input("Nhập Google Gemini API Key:", type="password")

st.sidebar.header("📘 Thông tin bài học")
subject = st.sidebar.text_input("Môn học:", "Ngữ văn")
grade = st.sidebar.text_input("Lớp:", "7")
lesson_name = st.sidebar.text_input("Tên bài học:", "Nhập tên bài học")
duration = st.sidebar.text_input("Thời lượng (tiết):", "1")
objectives = st.sidebar.text_area("Yêu cầu cần đạt (YCĐĐ):", "Nhập yêu cầu cần đạt")

# Hàm gọi AI Gemini
def generate_educational_content(key):
    client = genai.Client(api_key=key)
    
    prompt = f"""
    Bạn là một chuyên gia giáo dục phổ thông Việt Nam. Hãy soạn thảo nội dung bài học theo thông tin:
    - Môn: {subject}, Lớp: {grade}
    - Tên bài: {lesson_name} ({duration} tiết)
    - YCĐĐ: {objectives}

    Hãy trả về KẾT QUẢ ĐÚNG ĐỊNH DẠNG JSON với cấu trúc bắt buộc như sau:
    {{
        "giao_an": {{
            "muc_tieu": "Kiến thức, năng lực, phẩm chất...",
            "thiet_bi": "GV và HS...",
            "hoat_dong": [
                {{
                    "ten": "Hoạt động 1: Mở đầu",
                    "muc_tieu": "...",
                    "noi_dung": "...",
                    "san_pham": "...",
                    "to_chuc": "1. Chuyển giao... \\n2. Thực hiện... \\n3. Báo cáo... \\n4. Kết luận..."
                }}
            ]
        }},
        "slide": [
            {{"title": "Trang 1: Tên bài", "content": ["Nội dung 1", "Nội dung 2"]}}
        ],
        "de_thi_7991": {{
            "ma_tran": "Mô tả ma trận câu hỏi (Nhận biết 40%, Thông hiểu 30%, Vận dụng 20%, Vận dụng cao 10%)...",
            "cau_hoi": [
                {{"loai": "Trắc nghiệm 4 lựa chọn", "noi_dung": "Câu 1: ...", "dap_an": "A"}},
                {{"loai": "Đúng/Sai (CV 7991)", "noi_dung": "Câu 2: ...", "dap_an": "a-Đúng, b-Sai..."}}
            ]
        }}
    }}
    Chỉ trả về JSON thuần túy, không chứa văn bản dư thừa ngoài cấu trúc JSON.
    """
    
    response = client.models.generate_content(
        model='gemini-2.5-flash',
        contents=prompt
    )
    
    # Làm sạch chuỗi trả về để parse JSON
    text = response.text.replace("```json", "").replace("```", "").strip()
    return json.loads(text)

# Hàm xuất file Word (.docx) cho Giáo án
def export_docx(data):
    doc = Document()
    doc.add_heading(f"KẾ HOẠCH BÀI DẠY: {lesson_name.upper()}", level=1)
    
    giao_an = data.get("giao_an", {})
    doc.add_heading("I. MỤC TIÊU", level=2)
    doc.add_paragraph(giao_an.get("muc_tieu", ""))
    
    doc.add_heading("II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU", level=2)
    doc.add_paragraph(giao_an.get("thiet_bi", ""))
    
    doc.add_heading("III. TIẾN TRÌNH DẠY HỌC", level=2)
    for hd in giao_an.get("hoat_dong", []):
        doc.add_heading(hd.get("ten", ""), level=3)
        doc.add_paragraph(f"a) Mục tiêu: {hd.get('muc_tieu')}")
        doc.add_paragraph(f"b) Nội dung: {hd.get('noi_dung')}")
        doc.add_paragraph(f"c) Sản phẩm: {hd.get('san_pham')}")
        doc.add_paragraph(f"d) Tổ chức thực hiện:\n{hd.get('to_chuc')}")
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# Hàm xuất file PowerPoint (.pptx) cho Slide
def export_pptx(data):
    prs = Presentation()
    for slide_data in data.get("slide", []):
        slide_layout = prs.slide_layouts[1] # Bullet slide
        slide = prs.slides.add_slide(slide_layout)
        
        title = slide.shapes.title
        title.text = slide_data.get("title", "")
        
        body = slide.shapes.placeholders[1]
        tf = body.text_frame
        for bullet in slide_data.get("content", []):
            p = tf.add_paragraph()
            p.text = bullet
            
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()

# Nút xử lý
if st.button("🚀 Khởi tạo nội dung bài giảng & đề thi"):
    if not api_key:
        st.error("Vui lòng nhập Gemini API Key ở thanh bên trái!")
    else:
        with st.spinner("AI đang tạo Giáo án, Slide và Đề thi theo Công văn 7991..."):
            try:
                res_data = generate_educational_content(api_key)
                st.success("Tạo thành công!")
                
                tab1, tab2, tab3 = st.tabs(["📝 Kế hoạch bài dạy", "📊 Slide bài giảng", "📋 Đề kiểm tra (CV 7991)"])
                
                with tab1:
                    st.json(res_data.get("giao_an"))
                    docx_file = export_docx(res_data)
                    st.download_button("💾 Tải Giáo án (.docx)", docx_file, file_name="Giao_an_7991.docx")
                    
                with tab2:
                    st.json(res_data.get("slide"))
                    pptx_file = export_pptx(res_data)
                    st.download_button("💾 Tải Slide (.pptx)", pptx_file, file_name="Slide_Bai_Giang.pptx")
                    
                with tab3:
                    st.json(res_data.get("de_thi_7991"))
                    
            except Exception as e:
                st.error(f"Có lỗi xảy ra: {str(e)}")
